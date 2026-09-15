import calendar
import io
import re
import openpyxl
from openpyxl.styles import PatternFill
import pandas as pd
import streamlit as st
import xlrd

st.set_page_config(
    page_title="Portal Teknik LTPMS & STMPS", page_icon="⚙️", layout="wide"
)

st.title("⚙️ Portal Utama Maintenance (LTPMS & STMPS Generator)")
st.markdown("""
Pusat otomatisasi dokumen perencanaan pemeliharaan jangka panjang (**LTPMS**) dan break-down bulanan/harian (**STMPS**).
""")

# Sidebar Pengaturan Tahun & Bulan Target
st.sidebar.header("🔧 Pengaturan Target Dokumen")
target_year = st.sidebar.number_input(
    "Target Tahun", min_value=2025, max_value=2040, value=2027, step=1
)
selected_month = st.sidebar.selectbox(
    "Pilih Bulan untuk STMPS",
    options=list(range(1, 13)),
    format_func=lambda x: [
        "Januari",
        "Februari",
        "Maret",
        "April",
        "Mei",
        "Juni",
        "Juli",
        "Agustus",
        "September",
        "Oktober",
        "November",
        "Desember",
    ][x - 1],
    index=9,  # Default Oktober
)

y_n = target_year - 1
y_n1 = target_year - 2
y_n2 = target_year - 3
y_n3 = target_year - 4


def parse_history_ps(uploaded_file, year):
  ps_data = {}
  if not uploaded_file:
    return ps_data
  content = uploaded_file.read()
  uploaded_file.seek(0)

  if uploaded_file.name.endswith(".xlsx"):
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    for sname in wb.sheetnames:
      ws = wb[sname]
      for r in range(15, ws.max_row + 1):
        tag = str(ws.cell(r, 2).value or "").strip()
        name = str(ws.cell(r, 3).value or "").strip()
        if tag.endswith(".0"):
          tag = tag[:-2]
        clean_name = " ".join(name.upper().split())
        if not clean_name and not tag:
          continue

        months = []
        for c in range(4, 40):
          cell = ws.cell(r, c)
          if cell.fill and cell.fill.fill_type == "darkHorizontal":
            m = ((c - 4) // 3) + 1
            if m not in months:
              months.append(m)
        key = tag if tag else clean_name
        if months:
          ps_data[key] = sorted(months)
  else:
    wb = xlrd.open_workbook(file_contents=content, formatting_info=True)
    for sname in wb.sheet_names():
      if "sheet" in sname.lower() and sname.lower() != "1":
        continue
      sh = wb.sheet_by_name(sname)
      for r in range(14, sh.nrows):
        tag = str(sh.cell_value(r, 1)).strip() if sh.ncols > 1 else ""
        name = str(sh.cell_value(r, 2)).strip() if sh.ncols > 2 else ""
        if tag.endswith(".0"):
          tag = tag[:-2]
        clean_name = " ".join(name.upper().split())
        if not clean_name and not tag:
          continue

        months = []
        for c in range(3, min(39, sh.ncols)):
          xf = wb.xf_list[sh.cell_xf_index(r, c)]
          pat = xf.background.fill_pattern
          col = xf.background.pattern_colour_index
          is_ps = (
              col in (23, 24, 13)
              or pat == 5
              or (pat == 1 and col not in (8, 9, 64, 65))
          )
          if is_ps:
            m = ((c - 3) // 3) + 1
            if m not in months:
              months.append(m)
        key = tag if tag else clean_name
        if months:
          ps_data[key] = sorted(months)
  return ps_data


def process_ltpms(f_n3, f_n2, f_n1, f_n):
  ps_n3 = parse_history_ps(f_n3, y_n3) if f_n3 else {}
  ps_n2 = parse_history_ps(f_n2, y_n2) if f_n2 else {}
  ps_n1 = parse_history_ps(f_n1, y_n1) if f_n1 else {}

  content_master = f_n.read()
  f_n.seek(0)
  wb_master = openpyxl.load_workbook(
      io.BytesIO(content_master), data_only=False
  )

  if "1" in wb_master.sheetnames:
    wb_master["1"]["C7"].value = f":  {target_year}"

  ps_fill = PatternFill(
      fill_type="darkHorizontal", start_color="00000000", end_color="00000000"
  )
  pc_fill = PatternFill(
      fill_type="solid", start_color="00000000", end_color="00000000"
  )
  blank_fill = PatternFill(fill_type=None)

  scheduled_items = []

  for sname in wb_master.sheetnames:
    ws = wb_master[sname]
    for r in range(15, ws.max_row + 1):
      if r > ws.max_row - 8:
        row_txt = " ".join(
            [str(ws.cell(r, c).value or "") for c in range(1, 5)]
        ).upper()
        if "CHECK" in row_txt or "SERVICE" in row_txt or "NOTE" in row_txt:
          continue

      tag = str(ws.cell(r, 2).value or "").strip()
      name = str(ws.cell(r, 3).value or "").strip()
      rem = str(ws.cell(r, 40).value or "").strip()
      if tag.endswith(".0"):
        tag = tag[:-2]
      clean_name = " ".join(name.upper().split())
      if not clean_name and not tag:
        continue

      key = tag if tag else clean_name
      rem_upper = rem.upper()

      row_has_pc = False
      p_n = []
      for c in range(4, 40):
        cell = ws.cell(r, c)
        m = ((c - 4) // 3) + 1
        sc = (
            getattr(cell.fill.start_color, "index", None)
            if cell.fill
            else None
        )
        if (
            cell.fill
            and cell.fill.fill_type == "solid"
            and sc not in (9, "00000009", None)
        ):
          row_has_pc = True
        elif cell.fill and cell.fill.fill_type == "darkHorizontal":
          if m not in p_n:
            p_n.append(m)

      m_ps = re.search(r"PS\s*:\s*1\s*X\s*(\d+)\s*BLN", rem_upper)
      interval_ps = int(m_ps.group(1)) if m_ps else 12

      p_n3 = ps_n3.get(key, [])
      p_n2 = ps_n2.get(key, [])
      p_n1 = ps_n1.get(key, [])

      target_ps_months = []
      status_sched = ""

      if interval_ps == 12:
        target_ps_months = p_n if p_n else p_n1
        status_sched = "Rutin Tahunan (12 BLN)"
      elif interval_ps == 24:
        if p_n3 and p_n1:
          target_ps_months = p_n1 if p_n1 else p_n3
          status_sched = (
              f"Siklus Ganjil Asli ({y_n3} -> {y_n1} -> {target_year})"
          )
        elif p_n2:
          target_ps_months = []
          status_sched = f"Siklus Genap (Base {y_n2} -> Skip {target_year})"
        elif p_n1 or p_n3:
          target_ps_months = p_n1 if p_n1 else p_n3
          status_sched = f"Due Siklus 24 BLN dari {y_n1}"
      elif interval_ps == 36:
        if p_n2:
          target_ps_months = p_n2
          status_sched = f"Due Siklus 36 BLN dari {y_n2}"
        else:
          target_ps_months = []
          status_sched = f"Belum Jatuh Tempo ({interval_ps} BLN)"
      elif interval_ps == 48:
        if p_n3:
          target_ps_months = p_n3
          status_sched = f"Due Siklus 48 BLN dari {y_n3}"
        else:
          target_ps_months = []
          status_sched = f"Belum Jatuh Tempo ({interval_ps} BLN)"
      else:
        target_ps_months = p_n

      for m in p_n:
        if m not in target_ps_months:
          c_base = 4 + (m - 1) * 3
          fill_to_apply = pc_fill if row_has_pc else blank_fill
          for sub in range(3):
            ws.cell(r, c_base + sub).fill = fill_to_apply

      for m in target_ps_months:
        c_base = 4 + (m - 1) * 3
        for sub in range(3):
          ws.cell(r, c_base + sub).fill = ps_fill

      if target_ps_months:
        scheduled_items.append({
            "Sheet": sname,
            "Nama Mesin": name,
            "Interval": f"1x{interval_ps} BLN",
            "Keterangan": status_sched,
            "Bulan Servis": str(target_ps_months),
        })

  output_stream = io.BytesIO()
  wb_master.save(output_stream)
  output_stream.seek(0)
  return output_stream, scheduled_items


def is_ps_cell(c):
  return c.fill and c.fill.fill_type == "darkHorizontal"


def is_pc_cell(c):
  if not c.fill or c.fill.fill_type != "solid":
    return False
  sc = getattr(c.fill.start_color, "index", None)
  if sc in (9, "00000009", None):
    return False
  return True


def generate_stmps_from_template(
    f_ltpms, f_stmps_template, target_m, target_y, plant_name
):
  month_names = [
      "JANUARI",
      "FEBRUARI",
      "MARET",
      "APRIL",
      "MEI",
      "JUNI",
      "JULI",
      "AGUSTUS",
      "SEPTEMBER",
      "OKTOBER",
      "NOVEMBER",
      "DESEMBER",
  ]
  m_name = month_names[target_m - 1]

  # Deteksi Jumlah Hari & Tanggal Hari Minggu Secara Dinamis Berdasarkan Tahun & Bulan Target
  _, num_days = calendar.monthrange(target_y, target_m)
  sundays = [
      d
      for d in range(1, num_days + 1)
      if calendar.weekday(target_y, target_m, d) == 6
  ]

  # Susun Blok Hari Kerja Efektif (Terpisah Oleh Hari Minggu)
  week_working_days = {}
  w_idx = 0
  current_block = []

  for d in range(1, num_days + 1):
    if d in sundays:
      if current_block:
        week_working_days[w_idx] = current_block
        w_idx += 1
        current_block = []
    else:
      current_block.append(d)

  if current_block:
    week_working_days[w_idx] = current_block

  content_lt = f_ltpms.read()
  f_ltpms.seek(0)
  wb_lt = openpyxl.load_workbook(io.BytesIO(content_lt), data_only=False)

  content_st = f_stmps_template.read()
  f_stmps_template.seek(0)
  wb_st = openpyxl.load_workbook(io.BytesIO(content_st), data_only=False)

  ps_fill = PatternFill(
      fill_type="darkHorizontal", start_color="00000000", end_color="00000000"
  )
  pc_fill = PatternFill(
      fill_type="solid", start_color="00000000", end_color="00000000"
  )
  gray_fill = PatternFill(
      fill_type="solid", start_color="7F7F7F", end_color="7F7F7F"
  )
  blank_fill = PatternFill(fill_type=None)

  c_base_lt = 4 + (target_m - 1) * 3
  summary_rows = []

  for sname in wb_st.sheetnames:
    if sname not in wb_lt.sheetnames:
      continue
    ws_st = wb_st[sname]
    ws_lt = wb_lt[sname]

    # Update Header PERIODE
    ws_st["C8"].value = f":   {m_name} {target_y}"

    # Update Header Tanggal (1 s/d 31) -> Teks Selalu Muncul Terang!
    for d in range(1, 32):
      col_idx = 3 + d  # Kolom 4 = D (Tgl 1) s/d Kolom 34 = AH (Tgl 31)
      cell_hdr = ws_st.cell(13, col_idx)

      if d <= num_days:
        cell_hdr.value = float(d)
        if d in sundays:
          cell_hdr.fill = gray_fill
        else:
          cell_hdr.fill = blank_fill
      else:
        cell_hdr.value = ""
        cell_hdr.fill = blank_fill

    # Deteksi pekerjaan dari LTPMS
    lt_jobs = {}
    for r in range(15, ws_lt.max_row + 1):
      tag = str(ws_lt.cell(r, 2).value or "").strip()
      name = str(ws_lt.cell(r, 3).value or "").strip()
      if tag.endswith(".0"):
        tag = tag[:-2]
      clean_name = " ".join(name.upper().split())
      key = tag if tag else clean_name
      if not key:
        continue

      sub_cells = [ws_lt.cell(r, c_base_lt + i) for i in range(3)]
      has_ps = any(is_ps_cell(c) for c in sub_cells)
      has_pc = any(is_pc_cell(c) for c in sub_cells)

      if has_ps:
        lt_jobs[key] = "PS"
      elif has_pc:
        lt_jobs[key] = "PC"

    # Bersihkan area isi tabel (Baris 15 ke bawah): SEMUA JADI PUTIH BERSIH
    for r in range(15, ws_st.max_row + 1):
      for c in range(4, 35):
        ws_st.cell(r, c).fill = blank_fill

    # Pemetaan Arsir PC (Hitam) / PS (Garis) Sesuai Acuan Presisi
    for r in range(15, ws_st.max_row + 1):
      tag = str(ws_st.cell(r, 2).value or "").strip()
      name = str(ws_st.cell(r, 3).value or "").strip()
      if tag.endswith(".0"):
        tag = tag[:-2]
      clean_name = " ".join(name.upper().split())
      key = tag if tag else clean_name
      if not key or key not in lt_jobs:
        continue

      j_type = lt_jobs[key]
      target_fill = ps_fill if j_type == "PS" else pc_fill
      name_up = clean_name

      # Aturan Pemetaan Minggu Berdasarkan Urutan Mesin Pabrik (Sheet 1, 2, 3)
      if (
          "COMBUSTION" in name_up
          or "SCRUBBER" in name_up
          or "MAIN ROLLER NO.2" in name_up
          or "BLOWER ZONE RET" in name_up
          or "LONGITUDINAL CUTTING BRIDGE PG 01" in name_up
          or "EDGE TRIM ROLLER CONVEYOR" in name_up
          or "SNAPPING BRIDGE PG LEFT" in name_up
          or "ROLLER CONVEYOR PG POS.208" in name_up
          or "BLOWER FLOATING TABLE NO 3" in name_up
          or "BELT CONVEYOR PG 01" in name_up
      ):
        selected_w_idx = 0  # W1

      elif (
          "DRAIN GLASS" in name_up
          or "BLAST AIR" in name_up
          or "HOIST CRANE" in name_up
          or "BLOWER ZONE A/L" in name_up
          or "BLOWER ZONE B/L" in name_up
          or "BLOWER ZONE C/L" in name_up
          or "MAIN ROLLER NO.1" in name_up
          or "ROLLER ANNEALING" in name_up
          or "ROLLER CONVEYOR PG POS.201" in name_up
          or "LONGITUDINAL CUTTING BRIDGE PG 02" in name_up
          or "CROSS CUTTING BRIDGE PG 01" in name_up
          or "EDGE TRIM TOOLS" in name_up
          or "ROLLER CONV DROP SECTION" in name_up
          or "BELT CONVEYOR PG 02" in name_up
          or "CRUSHER PG 02" in name_up
      ):
        selected_w_idx = 1  # W2

      elif (
          "ROLLER TABLE" in name_up
          or "MEASURING BRIDGE" in name_up
          or "LONGITUDINAL CUTTING BRIDGE PG 03" in name_up
          or "CROSS CUTTING BRIDGE PG 02" in name_up
          or "SNAPPING BRIDGE PG RIGHT" in name_up
          or "ROLLER CONV CRUSHER INFEED" in name_up
      ):
        selected_w_idx = 2  # W3

      elif (
          "TURNING PLATFORM" in name_up
          or "BLOWER ZONE F2" in name_up
          or "LONGITUDINAL CUTTING BRIDGE PG 04" in name_up
          or "MAIN SNAPPING ROLLER" in name_up
          or "PLATE GLASS CRUSHER" in name_up
      ):
        selected_w_idx = 3  # W4

      elif (
          "ACCELERATION" in name_up
          or "MAIN DRIVE 02" in name_up
          or "BLOWER CHIPPING" in name_up
      ):
        selected_w_idx = 4  # W5

      else:
        selected_w_idx = 0

      # Gunakan blok minggu sesuai pembatas hari Minggu dinamis
      assigned_days = week_working_days.get(
          selected_w_idx, week_working_days.get(0, [1, 2])
      )

      for d in assigned_days:
        if d <= num_days:
          ws_st.cell(r, 3 + d).fill = target_fill

      summary_rows.append({
          "Sheet": sname,
          "Tag Equipment": tag,
          "Nama Mesin": name,
          "Jenis Pekerjaan": j_type,
          "Hari Kerja Eksekusi": (
              f"Tgl {assigned_days[0]} s/d {assigned_days[-1]} {m_name}"
          ),
      })

  output_stmps = io.BytesIO()
  wb_st.save(output_stmps)
  output_stmps.seek(0)
  return output_stmps, pd.DataFrame(summary_rows)


# ================= TAB BAGIAN ATAS =================
tab_float, tab_rolled, tab_stmps = st.tabs([
    "🏭 LTPMS Float 1",
    "🏭 LTPMS Rolled Glass",
    "📅 STMPS Generator (Format Asli Pabrik)",
])

# --- TAB 1: FLOAT 1 ---
with tab_float:
  st.subheader("📋 Generator LTPMS Float 1")
  col1, col2 = st.columns(2)
  with col1:
    fl_f_n3 = st.file_uploader(
        f"1. File Float 1 Tahun {y_n3} (.xls / .xlsx)",
        type=["xls", "xlsx"],
        key="fl_f_n3",
    )
    fl_f_n2 = st.file_uploader(
        f"2. File Float 1 Tahun {y_n2} (.xls / .xlsx)",
        type=["xls", "xlsx"],
        key="fl_f_n2",
    )
  with col2:
    fl_f_n1 = st.file_uploader(
        f"3. File Float 1 Tahun {y_n1} (.xls / .xlsx)",
        type=["xls", "xlsx"],
        key="fl_f_n1",
    )
    fl_f_n = st.file_uploader(
        f"4. File Float 1 Tahun {y_n} (Master Basis) (.xlsx)",
        type=["xlsx"],
        key="fl_f_n",
    )

  if st.button(
      "🚀 Proses Audit & Buat LTPMS Float 1", type="primary", key="btn_fl"
  ):
    if not fl_f_n:
      st.error(f"File Master Basis Tahun {y_n} untuk Float 1 wajib diunggah!")
    else:
      with st.spinner("Sedang memproses LTPMS Float 1..."):
        out_stream, sched_items = process_ltpms(
            fl_f_n3, fl_f_n2, fl_f_n1, fl_f_n
        )
        st.success(f"✅ LTPMS Float 1 {target_year} Berhasil Disusun!")
        st.download_button(
            label=f"📥 Unduh File Excel LTPMS Float 1 {target_year}",
            data=out_stream,
            file_name=f"LTPMS_FLOAT_1_{target_year}.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            key="dl_fl",
            type="primary",
        )
        st.dataframe(pd.DataFrame(sched_items), use_container_width=True)

# --- TAB 2: ROLLED GLASS ---
with tab_rolled:
  st.subheader("📋 Generator LTPMS Rolled Glass (Figur Glass)")
  col1, col2 = st.columns(2)
  with col1:
    rg_f_n3 = st.file_uploader(
        f"1. File Rolled Glass Tahun {y_n3} (.xls / .xlsx)",
        type=["xls", "xlsx"],
        key="rg_f_n3",
    )
    rg_f_n2 = st.file_uploader(
        f"2. File Rolled Glass Tahun {y_n2} (.xls / .xlsx)",
        type=["xls", "xlsx"],
        key="rg_f_n2",
    )
  with col2:
    rg_f_n1 = st.file_uploader(
        f"3. File Rolled Glass Tahun {y_n1} (.xls / .xlsx)",
        type=["xls", "xlsx"],
        key="rg_f_n1",
    )
    rg_f_n = st.file_uploader(
        f"4. File Rolled Glass Tahun {y_n} (Master Basis) (.xlsx)",
        type=["xlsx"],
        key="rg_n_basis",
    )

  if st.button(
      "🚀 Proses Audit & Buat LTPMS Rolled Glass", type="primary", key="btn_rg"
  ):
    if not rg_f_n:
      st.error(
          f"File Master Basis Tahun {y_n} untuk Rolled Glass wajib diunggah!"
      )
    else:
      with st.spinner("Sedang memproses LTPMS Rolled Glass..."):
        out_stream, sched_items = process_ltpms(
            rg_f_n3, rg_f_n2, rg_f_n1, rg_f_n
        )
        st.success(f"✅ LTPMS Rolled Glass {target_year} Berhasil Disusun!")
        st.download_button(
            label=f"📥 Unduh File Excel LTPMS Rolled Glass {target_year}",
            data=out_stream,
            file_name=f"LTPMS_ROLLED_GLASS_{target_year}.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            key="dl_rg",
            type="primary",
        )
        st.dataframe(pd.DataFrame(sched_items), use_container_width=True)

# --- TAB 3: STMPS GENERATOR FORMAT ASLI PABRIK ---
with tab_stmps:
  m_name_str = [
      "Januari",
      "Februari",
      "Maret",
      "April",
      "Mei",
      "Juni",
      "Juli",
      "Agustus",
      "September",
      "Oktober",
      "November",
      "Desember",
  ][selected_month - 1]
  st.subheader(f"📅 Generator STMPS Format Asli Pabrik (Bulan {m_name_str})")
  st.markdown(
      "Unggah berkas **LTPMS Hasil Audit** dan **Template Short Term Asli"
      " Pabrik (.xlsx)**."
  )

  col_s1, col_s2 = st.columns(2)
  with col_s1:
    plant_choice = st.radio(
        "Pilih Plant:",
        ["Rolled Glass (Figur Glass)", "Float 1"],
        horizontal=True,
    )
    f_ltpms_target = st.file_uploader(
        "1. Upload File LTPMS Hasil Audit (.xlsx) (Wajib)",
        type=["xlsx"],
        key="f_stmps_ltpms",
    )
  with col_s2:
    f_stmps_template = st.file_uploader(
        "2. Upload Template Short Term Asli Pabrik (.xlsx) (Wajib)",
        type=["xlsx"],
        key="f_stmps_tpl",
    )

  if st.button(
      f"⚡ Generate STMPS {plant_choice} Bulan {m_name_str}", type="primary"
  ):
    if not f_ltpms_target or not f_stmps_template:
      st.error(
          "Silakan unggah kedua file (LTPMS Hasil Audit & Template Short Term"
          " Asli) terlebih dahulu!"
      )
    else:
      with st.spinner(f"Sedang memproses STMPS bulan {m_name_str}..."):
        stmps_out, df_sum = generate_stmps_from_template(
            f_ltpms_target,
            f_stmps_template,
            selected_month,
            target_year,
            plant_choice,
        )
        m_name_up = m_name_str.upper()
        p_label = (
            "PIGURE_GLASS"
            if "ROLLED" in plant_choice.upper()
            or "PIGUR" in plant_choice.upper()
            else "FLOAT_1"
        )

        st.success(
            f"✅ STMPS {plant_choice} {m_name_up} {target_year} Berhasil"
            " Disusun!"
        )
        st.download_button(
            label=(
                f"📥 Unduh File Excel SHORT TERM {p_label} {m_name_up}"
                f" {target_year}"
            ),
            data=stmps_out,
            file_name=(
                f"SHORT_TERM_{p_label}_{m_name_up}_{target_year}.xlsx"
            ),
            mime=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            type="primary",
        )
        st.dataframe(df_sum, use_container_width=True)
